"""
Look-through das carteiras de fundos (fonte: Mais Retorno, dados CVM/CDA).
Abre recursivamente toda cota de fundo que tenha carteira aberta; a que nao tiver vira ativo final
(categoria FIDC / FIAGRO / Cotas de Fundos). % final = produto dos % ao longo da cadeia.
Saida:
  lookthrough_peers.xlsx
    graficos   % do PL por categoria ao longo do tempo, um grafico por peer
    evolucao   dados dos graficos (mes x categoria)
    top10      top 10 ativos por categoria (ultimo mes de cada peer), entre todos os peers
    <peer>     mes > categoria > ativos (linhas agrupadas, clique no +)
  lookthrough_peers_dados.csv   tabela plana (uma linha por caminho ate o ativo)
"""
import json
import re
import sqlite3
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from pathlib import Path

import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import AreaChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.indexed_list import IndexedList

PEERS = {                       # nome: link do Mais Retorno OU CNPJ
    "DUAL": "https://maisretorno.com/fundo/itau-dual-private-markets-multimercado-cp-fif-rl",
    "AZ_ALTRO": "22.100.009/0001-07",
    "SPARTA_TOP": "https://maisretorno.com/fundo/sparta-top-fic-fif-rf-cp-lp-rl-1",
    "XP_CE120": "22.003.930/0001-31",
    "CAPITANIA_P45": "https://maisretorno.com/fundo/capitania-premium-45-fic-fi-rf-cp-lp",
}
DESDE = None                    # ex "2025-01" para limitar; None = desde a carteira mais antiga
OUT = "lookthrough_peers.xlsx"
THREADS = 3                     # o site bloqueia (403) se for rapido demais: a velocidade vem do cache, nao de mais threads

API = "https://data.maisretorno.com/mr-data/v4/fund"
H = {"User-Agent": "Mozilla/5.0", "Referer": "https://maisretorno.com"}

# cache num unico arquivo SQLite (milhares de .json soltos ficam lentos, principalmente no OneDrive)
DB = sqlite3.connect("cache_maisretorno.sqlite", check_same_thread=False)
DB.execute("CREATE TABLE IF NOT EXISTS c (k TEXT PRIMARY KEY, v TEXT)")
MEM, LOCK, LOCAL = {}, threading.Lock(), threading.local()


class Falha(Exception):              # resposta nao definitiva (bloqueio/erro): nao vira "carteira vazia"
    pass


def get(path):
    # carteiras mensais nao mudam: cache permanente; lista de meses: cache do dia
    key = path + (f"@{date.today()}" if "available" in path else "")
    with LOCK:
        if key in MEM:
            return MEM[key]
        row = DB.execute("SELECT v FROM c WHERE k=?", (key,)).fetchone()
    if row:
        data = json.loads(row[0])
    else:
        if not hasattr(LOCAL, "s"):                    # uma sessao HTTP por thread: reaproveita a conexao
            LOCAL.s = requests.Session()
            LOCAL.s.headers.update(H)
        r = None
        for espera in (0.2, 5, 20, 60, 120):           # 403/429/5xx = site pedindo calma: espera e tenta de novo
            time.sleep(espera)
            try:
                r = LOCAL.s.get(f"{API}/{path}", timeout=60)
                if r.status_code in (200, 404):
                    break
            except requests.RequestException:
                r = None
        if r is None or r.status_code not in (200, 404):
            raise Falha(f"{path} -> {None if r is None else r.status_code}")
        data = r.json() if r.ok else None
        if True:                                        # so respostas definitivas (200/404) chegam aqui e vao pro cache
            with LOCK:
                DB.execute("INSERT OR REPLACE INTO c VALUES (?, ?)", (key, json.dumps(data, ensure_ascii=False)))
    with LOCK:
        MEM[key] = data
    return data

def cnpj_of(x):
    d = re.sub(r"\D", "", x)
    if "maisretorno" not in x and len(d) == 14:
        return d
    m = re.search(r"id=(\d{14})", x)
    if m:
        return m.group(1)
    html = requests.get(x, headers=H, timeout=60).text
    return re.sub(r"\D", "", re.search(r"CNPJ\D{0,30}(\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2})", html).group(1))


def months(cnpj):
    return sorted(m["date"][:7] for m in get(f"available-wallets/{cnpj}") or [] if m["open"])


def ref_month(cnpj, ym):             # mesma data, ou a carteira aberta mais recente antes dela
    ok = [m for m in months(cnpj) if m <= ym]
    return ok[-1] if ok else None


def codigo(a):
    m = re.search(r"id=(\d{14})", a.get("canonical_url") or "")
    if m:
        return m.group(1)
    for pat in (r"\b(BR[A-Z0-9]{9}\d)\b", r"\b((?:CRI|CRA):\S+)", r"^([A-Z0-9]{5,6}) - ", r" - ([A-Z]{4}[A-Z0-9]{1,2})$"):
        m = re.search(pat, a["nicename"] or "")
        if m:
            return m.group(1)
    return ""


def fund_cat(name):
    n = name.upper()
    return ("FIDC" if "FIDC" in n else "FIAGRO" if "FIAGRO" in n or "CADEIAS PRODUTIVAS" in n
            else "Cotas de Fundos")


def explode(cnpj, ym, peso, caminho, out):
    w = get(f"wallet-detail/{cnpj}?year={ym[:4]}&month={ym[5:]}") or {}
    for c in w.get("classes", []):
        for a in c["assets"]:
            p = peso * (a.get("percentage") or 0) / 100
            filho = re.search(r"id=(\d{14})", a.get("canonical_url") or "")
            cat = c["nicename"]
            if cat == "Cotas de Fundos":
                f = filho.group(1) if filho else None
                ref = ref_month(f, ym) if f and f not in [x[0] for x in caminho] else None
                if ref:
                    explode(f, ref, p, caminho + [(f, a["nicename"])], out)
                    continue
                cat = fund_cat(a["nicename"])
            out.append(dict(categoria=cat, ativo=a["nicename"], codigo=codigo(a), perc_pl=p * 100,
                            veiculo=caminho[-1][1], cnpj_veiculo=caminho[-1][0],
                            caminho=" > ".join(x[1] for x in caminho), mes_carteira=ym))


def job(args):
    nome, cnpj, ym = args
    out = []
    try:
        explode(cnpj, ym, 1.0, [(cnpj, nome)], out)
    except Falha:
        return None                       # carteira incompleta: tenta de novo no fim
    return [dict(fundo=nome, mes=ym, **r) for r in out]


CNPJ, jobs = {}, []
for n, x in PEERS.items():
    try:
        CNPJ[n] = cnpj_of(x)
        jobs += [(n, CNPJ[n], ym) for ym in months(CNPJ[n]) if not DESDE or ym >= DESDE]
    except (Falha, requests.RequestException, AttributeError) as e:
        print(f"ATENCAO: {n} pulado ({e}); confira o link/CNPJ ou rode de novo mais tarde")
print(f"{len(jobs)} carteiras (fundo x mes) para abrir")
rows, falhas = [], []
with ThreadPoolExecutor(THREADS) as ex:
    for i, r in enumerate(ex.map(job, jobs), 1):
        if r is None:
            falhas.append(jobs[i - 1])
        else:
            rows += r
        if i % 25 == 0 or i == len(jobs):
            print(f"  {i}/{len(jobs)}  ({jobs[i - 1][0]} {jobs[i - 1][2]})")
            with LOCK:
                DB.commit()                      # grava o cache aos poucos: se cair, nao perde o que ja baixou
with LOCK:
    DB.commit()
if falhas:                                # segunda passada, devagar, so no que falhou
    print(f"{len(falhas)} carteiras falharam (site bloqueou); tentando de novo em 60s, uma por vez...")
    time.sleep(60)
    falhas, refazer = [], falhas
    for j in refazer:
        r = job(j)
        if r is None:
            falhas.append(j)
        else:
            rows += r
    with LOCK:
        DB.commit()
if falhas:
    print("ATENCAO: ficaram de fora (rode de novo mais tarde; o cache guarda o resto):")
    for n, _, ym in falhas:
        print("   ", n, ym)
df = pd.DataFrame(rows)


def limpa(v):                       # conserta acentos quebrados na fonte (Ã‡ -> Ç) e tira caracteres que o Excel recusa
    if not isinstance(v, str):
        return v
    if "Ã" in v:
        try:
            v = v.encode("latin1").decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            pass
    return ILLEGAL_CHARACTERS_RE.sub("", v)


for c in df.select_dtypes("object"):
    df[c] = df[c].map(limpa)

# ---------------------------------------------------------------- checks
soma = df.groupby(["fundo", "mes"]).perc_pl.sum()
print("soma do %PL por carteira: min", round(soma.min(), 4), "max", round(soma.max(), 4))
if (soma < 99).any():
    print("carteiras somando < 99% (dado da fonte, ex. cota sem %):", soma[soma < 99].round(2).to_dict())
cruz = df[df.cnpj_veiculo.isin(CNPJ.values()) & (df.cnpj_veiculo != df.fundo.map(CNPJ))]
print("peers investindo em outro peer:", "nenhum" if cruz.empty else cruz[["fundo", "veiculo"]].drop_duplicates().to_string())

# ---------------------------------------------------------------- Excel
B = Font(name="Arial", size=10, bold=True)
N = Font(name="Arial", size=10)
HF = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F3864")
MFILL = PatternFill("solid", fgColor="D9E1F2")
PCT = '0.0000"%";(0.0000"%")'
PCT2 = '0.00"%";(0.00"%")'


def put(ws, vals, font=None, fill=None, fmt=None):
    # ws.max_row / ws[row] recontam a planilha inteira a cada chamada (lento); _current_row e O(1)
    ws.append(vals)
    r = ws._current_row
    for j in range(1, len(vals) + 1):
        c = ws.cell(r, j)
        if font:
            c.font = font
        if fill:
            c.fill = fill
        if fmt and j in fmt:
            c.number_format = fmt[j]
    return r


def header(ws, cols, widths):
    put(ws, cols, HF, FILL)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w


wb = Workbook()
wb._fonts = IndexedList([N])                 # Arial 10 como fonte padrao: nao precisa formatar celula por celula
wb.remove(wb.active)

# 1) uma aba por peer: mes > categoria > ativos agrupados
for nome, g in df.groupby("fundo", sort=False):
    ws = wb.create_sheet(nome[:31])
    header(ws, ["Categoria / Ativo", "Codigo", "% PL " + nome, "% da categoria", "Veiculo", "Mes carteira"],
           (70, 18, 14, 14, 50, 12))
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.freeze_panes = "A2"
    for ym, gm in sorted(g.groupby("mes"), reverse=True):
        put(ws, [ym, None, gm.perc_pl.sum()], B, MFILL, {3: PCT})
        ats = (gm.groupby(["categoria", "ativo", "codigo"])
               .agg(perc_pl=("perc_pl", "sum"), veiculo=("veiculo", "first"), mes_carteira=("mes_carteira", "first")))
        for cat, tot in gm.groupby("categoria").perc_pl.sum().sort_values(ascending=False).items():
            r = put(ws, ["   " + cat, None, tot], B, None, {3: PCT})
            ws.row_dimensions[r].outline_level = 1
            for (ativo, cod), a in ats.loc[cat].sort_values("perc_pl", ascending=False).iterrows():
                last = put(ws, ["      " + ativo, cod, a.perc_pl, a.perc_pl / tot * 100 if tot else None, a.veiculo,
                                a.mes_carteira], fmt={3: PCT, 4: PCT})
            ws.row_dimensions.group(r + 1, last, outline_level=2, hidden=True)
# 2) evolucao por categoria + graficos (7 maiores categorias no total + Outros; cor fixa por categoria)
CURTO = {"Depósitos a prazo e outros títulos de IF": "Depósitos a prazo / IF",
         "Outros valores mobiliários registrados na CVM objeto de oferta pública": "Outros VM (CVM)",
         "Títulos ligados ao agronegócio": "Títulos agro", "Títulos de Crédito Privado": "Crédito privado (CCB/NC)"}
CORES = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "4A3AA7"]
OUTROS_COR = "898781"
cat_ = df.categoria.map(lambda c: CURTO.get(c, c))
CONTABIL = r"(?i)pagar|receber|obriga|termo|disponibilidade|exigibilidade|swap"   # linhas contabeis vao para Outros
TOP = (df.assign(c=cat_)[~cat_.str.contains(CONTABIL)].groupby(["fundo", "mes", "c"]).perc_pl.sum()
       .groupby("c").mean().sort_values(ascending=False).index[:7].tolist())
SERIES = TOP + ["Outros"]
ev = df.assign(c=cat_.where(cat_.isin(TOP), "Outros")).pivot_table(
    index=["fundo", "mes"], columns="c", values="perc_pl", aggfunc="sum", fill_value=0)[SERIES]

ws = wb.create_sheet("evolucao", 0)
gr = wb.create_sheet("graficos", 0)
gr.append(["Composicao por categoria (% do PL, look-through). Dados em 'evolucao'."])
gr["A1"].font = B
for k, nome in enumerate(p for p in PEERS if p in ev.index.get_level_values(0)):
    t = ev.loc[nome]
    row0 = put(ws, [nome], B)
    put(ws, ["mes"] + SERIES, HF, FILL)
    for ym, r in t.iterrows():
        last = put(ws, [ym] + [float(v) for v in r], fmt={j: PCT2 for j in range(2, len(SERIES) + 2)})
    ch = AreaChart()
    ch.grouping = "stacked"
    ch.title = f"{nome} — % do PL por categoria"
    ch.y_axis.title = "% do PL"
    ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 100
    ch.y_axis.majorGridlines = None
    ch.x_axis.tickLblSkip = 12
    ch.x_axis.delete = ch.y_axis.delete = False
    ch.add_data(Reference(ws, min_col=2, max_col=1 + len(SERIES), min_row=row0 + 1, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=row0 + 2, max_row=last))
    for s, cor in zip(ch.series, CORES + [OUTROS_COR]):
        s.graphicalProperties.solidFill = cor
        s.graphicalProperties.line.solidFill = "FCFCFB"
    ch.legend.position = "r"
    ch.width, ch.height = 26, 9
    gr.add_chart(ch, f"A{3 + k * 19}")
    ws.append([])
    ws.append([])
ws.column_dimensions["A"].width = 12
for i in range(len(SERIES)):
    ws.column_dimensions[chr(66 + i)].width = 16

# 3) top 10 ativos por categoria, ultimo mes de cada peer, entre todos
ult = df[df.mes == df.fundo.map(df.groupby("fundo").mes.max())]
agg = (ult.groupby(["categoria", "fundo", "mes", "ativo", "codigo"]).perc_pl.sum().reset_index()
       .sort_values(["categoria", "perc_pl"], ascending=[True, False]))
top10 = agg.groupby("categoria").head(10)
top10.insert(1, "rank", top10.groupby("categoria").cumcount() + 1)
ws = wb.create_sheet("top10", 2)
header(ws, ["categoria", "rank", "fundo", "mes", "ativo", "codigo", "% PL do fundo"], (34, 6, 16, 9, 70, 18, 14))
for r in top10.itertuples(index=False):
    put(ws, list(r), fmt={7: PCT})

wb.save(OUT)
df.to_csv(OUT.replace(".xlsx", "_dados.csv"), index=False, sep=";", decimal=",", encoding="utf-8-sig")
print("salvo", OUT, "e", OUT.replace(".xlsx", "_dados.csv"), len(df), "linhas")